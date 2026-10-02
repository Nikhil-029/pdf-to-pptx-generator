import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from app.schemas.schemas import PresentationOutline, SlideOutlineItem, DeckTheme

class PptxRendererService:
    @staticmethod
    def hex_to_rgb(hex_str: str) -> RGBColor:
        """Converts a hex string like '#1E2761' to python-pptx RGBColor."""
        hex_str = hex_str.lstrip('#')
        if len(hex_str) != 6:
            # Fallback to black if invalid
            return RGBColor(0, 0, 0)
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

    @staticmethod
    def apply_slide_background(slide, color: RGBColor):
        """Applies a solid background color to a slide."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    @staticmethod
    def format_text_frame(text_frame, text: str, font_size: Pt, font_color: RGBColor, is_bold: bool = False, alignment=PP_ALIGN.LEFT):
        """Standardizes text formatting across all shapes."""
        text_frame.text = text
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.font.size = font_size
        p.font.color.rgb = font_color
        p.font.bold = is_bold
        p.font.name = "Calibri" # A safe, cross-platform font (avoiding Aptos per rules)
        p.alignment = alignment

    @staticmethod
    def generate_presentation(outline: PresentationOutline, output_filename: str, slide_images: dict = None) -> str:
        """
        Translates the Pydantic outline into a physical .pptx file.
        Accepts an optional slide_images dictionary mapping slide_number -> local_image_path.
        """
        if slide_images is None:
            slide_images = {}

        prs = Presentation()
        
        # PPTX default slide dimensions (16:9 widescreen)
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        theme = outline.deck_theme
        primary_rgb = PptxRendererService.hex_to_rgb(theme.primary_color)
        secondary_rgb = PptxRendererService.hex_to_rgb(theme.secondary_color)
        accent_rgb = PptxRendererService.hex_to_rgb(theme.accent_color)
        
        # We use a completely blank layout (index 6) so we can draw our own custom layouts
        blank_slide_layout = prs.slide_layouts[6] 

        for slide_data in outline.slides:
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Extract the downloaded image path for this specific slide (if available)
            local_image_path = slide_images.get(slide_data.slide_number)

            # --- Dark/Light Sandwich Logic ---
            # Titles and Section Headers get dark backgrounds. Content gets light.
            is_dark_slide = slide_data.layout in ["Title", "SectionHeader"]
            
            bg_color = primary_rgb if is_dark_slide else accent_rgb
            text_main_color = accent_rgb if is_dark_slide else primary_rgb
            text_muted_color = secondary_rgb if is_dark_slide else PptxRendererService.hex_to_rgb("#555555")

            PptxRendererService.apply_slide_background(slide, bg_color)

            # --- Layout Routing ---
            if slide_data.layout in ["Title", "SectionHeader"]:
                PptxRendererService._draw_title_layout(slide, slide_data, text_main_color, text_muted_color)
            elif slide_data.layout == "TwoColumn":
                PptxRendererService._draw_two_column_layout(slide, slide_data, text_main_color, text_muted_color, local_image_path)
            elif slide_data.layout == "HalfBleedImage":
                PptxRendererService._draw_half_bleed_layout(slide, slide_data, text_main_color, text_muted_color, secondary_rgb, local_image_path)
            else:
                # Default fallback for Grid2x2, IconTextRows, or Summary
                PptxRendererService._draw_standard_content_layout(slide, slide_data, text_main_color)

            # Add Speaker Notes
            if slide_data.speaker_notes:
                notes_slide = slide.notes_slide
                notes_text_frame = notes_slide.notes_text_frame
                notes_text_frame.text = slide_data.speaker_notes

        prs.save(output_filename)
        return output_filename

    # --- Specific Layout Drawers ---

    @staticmethod
    def _draw_title_layout(slide, data: SlideOutlineItem, main_color: RGBColor, sub_color: RGBColor):
        # Title Box
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
        PptxRendererService.format_text_frame(
            title_box.text_frame, data.title, Pt(44), main_color, is_bold=True, alignment=PP_ALIGN.CENTER
        )
        
        # Subtitle / Body Box (if any)
        if data.body_content:
            body_text = " | ".join(data.body_content) # Flatten bullets for title slide
            body_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(1))
            PptxRendererService.format_text_frame(
                body_box.text_frame, body_text, Pt(20), sub_color, alignment=PP_ALIGN.CENTER
            )

    @staticmethod
    def _draw_two_column_layout(slide, data: SlideOutlineItem, main_color: RGBColor, sub_color: RGBColor, image_path: str = None):
        # Header
        header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
        PptxRendererService.format_text_frame(header_box.text_frame, data.title, Pt(32), main_color, is_bold=True)

        # Left Column (Text)
        left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(5.5), Inches(5))
        tf = left_box.text_frame
        tf.word_wrap = True
        for i, point in enumerate(data.body_content):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"• {point}"
            p.font.size = Pt(18)
            p.font.color.rgb = main_color
            p.space_after = Pt(14) # Breathing room between bullets

        # Right Column (Image or Placeholder)
        if image_path and os.path.exists(image_path):
            slide.shapes.add_picture(
                image_path, 
                Inches(6.5), Inches(1.8), width=Inches(6), height=Inches(4.5)
            )
        else:
            right_box = slide.shapes.add_shape(
                1, Inches(6.5), Inches(1.8), Inches(6), Inches(4.5) # 1 is AutoShapeType for Rectangle
            )
            right_box.fill.solid()
            right_box.fill.fore_color.rgb = PptxRendererService.hex_to_rgb("#EEEEEE") # Light grey placeholder
            right_box.line.fill.background()
            
            placeholder_tf = right_box.text_frame
            PptxRendererService.format_text_frame(
                placeholder_tf, f"[ Visual Requirement: {data.visual_instruction} ]", 
                Pt(14), sub_color, alignment=PP_ALIGN.CENTER
            )

    @staticmethod
    def _draw_half_bleed_layout(slide, data: SlideOutlineItem, main_color: RGBColor, sub_color: RGBColor, accent_bg: RGBColor, image_path: str = None):
        # Left side is a massive image or color block
        if image_path and os.path.exists(image_path):
            slide.shapes.add_picture(
                image_path, 
                Inches(0), Inches(0), width=Inches(5.5), height=Inches(7.5)
            )
        else:
            left_bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(5.5), Inches(7.5))
            left_bg.fill.solid()
            left_bg.fill.fore_color.rgb = accent_bg
            left_bg.line.fill.background()
            
            img_instruction_tf = left_bg.text_frame
            PptxRendererService.format_text_frame(
                img_instruction_tf, f"[ Full Height Image Placeholder ]\n\n{data.visual_instruction}", 
                Pt(16), PptxRendererService.hex_to_rgb("#FFFFFF"), alignment=PP_ALIGN.CENTER
            )

        # Right side is Content
        header_box = slide.shapes.add_textbox(Inches(6.0), Inches(1.5), Inches(6.5), Inches(1))
        PptxRendererService.format_text_frame(header_box.text_frame, data.title, Pt(36), main_color, is_bold=True)

        body_box = slide.shapes.add_textbox(Inches(6.0), Inches(3), Inches(6.5), Inches(4))
        tf = body_box.text_frame
        tf.word_wrap = True
        for i, point in enumerate(data.body_content):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = point
            p.font.size = Pt(18)
            p.font.color.rgb = main_color
            p.space_after = Pt(14)

    @staticmethod
    def _draw_standard_content_layout(slide, data: SlideOutlineItem, main_color: RGBColor):
        # Simple Header + Bullets
        header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
        PptxRendererService.format_text_frame(header_box.text_frame, data.title, Pt(32), main_color, is_bold=True)

        body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(11.5), Inches(5))
        tf = body_box.text_frame
        tf.word_wrap = True
        for i, point in enumerate(data.body_content):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"• {point}"
            p.font.size = Pt(20)
            p.font.color.rgb = main_color
            p.space_after = Pt(16)